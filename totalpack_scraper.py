#!/usr/bin/env python3
"""Scrape selected TotalPack SKUs and fill the supplied Temu template."""

from __future__ import annotations

import argparse
import csv
import html
import json
import logging
import math
import re
import shutil
import sys
import time
from dataclasses import dataclass, field
from datetime import datetime, timezone
from html.parser import HTMLParser
from pathlib import Path
from typing import Any, Iterable
from urllib.error import HTTPError, URLError
from urllib.parse import urlencode
from urllib.request import Request, urlopen

from openpyxl import load_workbook


LOG = logging.getLogger("totalpack")
STORE_API = "https://totalpack.bg/wp-json/wc/store/v1/products"
DEFAULT_TEMPLATE = "TEMU_ALL-PURPOSE-LABELS_ALUMINUM-FOIL_BOXES_BOXES_CLOTH-FACE-MASKS_CONTAINERS_CONTAINERS_DISPOSABLE-MEDIC.xlsx"
DEFAULT_SKU_FILE = "totalpack_product_skus.xlsx"


CATEGORY_NAMES = {
    1408: "Office Products / Office & School Supplies / Labels, Indexes & Stamps / Labels & Stickers / All-Purpose Labels",
    10628: "Home & Kitchen / Kitchen & Dining / Dining & Entertaining / Flatware / Flatware Sets / Flatware Sets",
    53765: "Industrial & Scientific / Professional Medical Supplies / Apparel / Face Masks & Shields / Disposable Medical Masks",
    9030: "Industrial & Scientific / Retail Store Fixtures & Equipment / Point-of-Sale (POS) Equipment & Accessories / POS & Register Rolls",
    843: "Office Products / Office & School Supplies / Envelopes, Mailers & Shipping Supplies / Packing Materials / Packing Tape",
    11085: "Home & Kitchen / Kitchen & Dining / Food Service Equipment & Supplies / Disposables / Take Out Containers / Boxes",
    31121: "Clothing, Shoes & Jewelry / Shoe, Jewelry & Watch Accessories / Shoe Care & Accessories / Shoe Covers & Overshoes / Shoe Covers",
    11091: "Home & Kitchen / Kitchen & Dining / Food Service Equipment & Supplies / Disposables / Take Out Containers / Pizza Boxes",
    12950: "Home & Kitchen / Cleaning Supplies / Trash Bags",
    17188: "Health & Household / Household Supplies / Paper & Plastic / Disposable Food Storage / Food Storage Bags",
    17189: "Health & Household / Household Supplies / Paper & Plastic / Disposable Food Storage / Plastic Wrap",
    17190: "Health & Household / Household Supplies / Paper & Plastic / Disposable Food Storage / Aluminum Foil",
    10088: "Home & Kitchen / Kitchen & Dining / Kitchen Utensils & Gadgets / Straws / Disposable Straws",
    6774: "Industrial & Scientific / Occupational Health & Safety Products / Personal Protective Equipment / Hand & Arm Protection / Lab, Safety & Work Gloves / Disposable Gloves / Non-Sterile Gloves",
}


@dataclass
class SourceSku:
    sku: str
    name: str = ""
    weight_kg: float = 0.0
    length_cm: float = 0.0
    width_cm: float = 0.0
    height_cm: float = 0.0
    url: str = ""
    source_category: str = ""


@dataclass
class TemuRow:
    sku: str
    name: str
    description: str
    bullets: list[str]
    images: list[str]
    url: str
    category_id: int
    price: float
    list_price: float | None
    quantity: int
    weight_g: int
    length_cm: float
    width_cm: float
    height_cm: float
    dimensions_text: str
    pack_count: int
    properties: dict[str, Any] = field(default_factory=dict)
    notes: list[str] = field(default_factory=list)


class TextExtractor(HTMLParser):
    BREAK_TAGS = {"br", "p", "li", "div", "h1", "h2", "h3", "tr"}

    def __init__(self) -> None:
        super().__init__()
        self.parts: list[str] = []

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag.lower() in self.BREAK_TAGS:
            self.parts.append("\n")
        if tag.lower() == "li":
            self.parts.append("• ")

    def handle_endtag(self, tag: str) -> None:
        if tag.lower() in self.BREAK_TAGS:
            self.parts.append("\n")

    def handle_data(self, data: str) -> None:
        self.parts.append(data)


def clean_html(value: str | None) -> str:
    parser = TextExtractor()
    parser.feed(value or "")
    text = html.unescape("".join(parser.parts)).replace("\xa0", " ")
    lines = [re.sub(r"\s+", " ", line).strip() for line in text.splitlines()]
    return "\n".join(line for line in lines if line).strip()


def normalize_sku(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    text = str(value).strip()
    return text[:-2] if re.fullmatch(r"\d+\.0", text) else text


def number(value: Any, default: float = 0.0) -> float:
    if value is None or value == "":
        return default
    try:
        result = float(str(value).replace(",", "."))
        return result if math.isfinite(result) else default
    except (TypeError, ValueError):
        return default


def read_sku_file(path: Path) -> list[SourceSku]:
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    header_row = None
    sku_col = 1
    for row in range(1, min(ws.max_row, 20) + 1):
        for col in range(1, min(ws.max_column, 20) + 1):
            if str(ws.cell(row, col).value or "").strip().upper() == "SKU":
                header_row, sku_col = row, col
                break
        if header_row:
            break
    if not header_row:
        raise ValueError(f"No SKU header found in {path}")

    rows: list[SourceSku] = []
    seen: set[str] = set()
    for row in range(header_row + 1, ws.max_row + 1):
        sku = normalize_sku(ws.cell(row, sku_col).value)
        if not sku or sku in seen:
            continue
        seen.add(sku)
        # The supplied TotalPack export has a stable 13-column structure:
        # SKU, name, slug, short desc, description, price, weight kg,
        # length, width, height, image data, URL, category path.
        rows.append(
            SourceSku(
                sku=sku,
                name=str(ws.cell(row, 2).value or "").strip(),
                weight_kg=number(ws.cell(row, 7).value),
                length_cm=number(ws.cell(row, 8).value),
                width_cm=number(ws.cell(row, 9).value),
                height_cm=number(ws.cell(row, 10).value),
                url=str(ws.cell(row, 12).value or "").strip(),
                source_category=str(ws.cell(row, 13).value or "").strip(),
            )
        )
    wb.close()
    return rows


def http_json(url: str, timeout: int, retries: int) -> Any:
    last_error: Exception | None = None
    for attempt in range(1, retries + 1):
        try:
            request = Request(
                url,
                headers={
                    "User-Agent": "Mozilla/5.0 (compatible; TotalPackTemuScraper/1.0)",
                    "Accept": "application/json",
                    "Accept-Language": "bg-BG,bg;q=0.9,en;q=0.8",
                },
            )
            with urlopen(request, timeout=timeout) as response:
                return json.loads(response.read().decode("utf-8"))
        except (HTTPError, URLError, TimeoutError, json.JSONDecodeError) as exc:
            last_error = exc
            if attempt < retries:
                delay = min(2 ** (attempt - 1), 8)
                LOG.warning("Request failed (%s/%s): %s; retrying in %ss", attempt, retries, exc, delay)
                time.sleep(delay)
    raise RuntimeError(f"Request failed after {retries} attempts: {url}: {last_error}")


def chunks(values: list[str], size: int) -> Iterable[list[str]]:
    for start in range(0, len(values), size):
        yield values[start : start + size]


def fetch_products(skus: list[str], timeout: int, retries: int, fixture: Path | None) -> dict[str, dict[str, Any]]:
    if fixture:
        data = json.loads(fixture.read_text(encoding="utf-8"))
        return {normalize_sku(item.get("sku")): item for item in data if normalize_sku(item.get("sku"))}

    products: dict[str, dict[str, Any]] = {}
    for group in chunks(skus, 50):
        query = urlencode({"sku": ",".join(group), "per_page": 100})
        data = http_json(f"{STORE_API}?{query}", timeout, retries)
        if not isinstance(data, list):
            raise RuntimeError("Unexpected response from the WooCommerce Store API")
        for item in data:
            sku = normalize_sku(item.get("sku"))
            if sku:
                products[sku] = item
        LOG.info("Fetched %s/%s requested SKUs", len(products), len(skus))
    return products


def pick_category(product: dict[str, Any], source: SourceSku) -> int:
    slugs = {str(c.get("slug", "")).lower() for c in product.get("categories", [])}
    text = " ".join(
        [
            clean_html(product.get("name")),
            clean_html(product.get("short_description")),
            source.source_category,
            " ".join(slugs),
        ]
    ).lower()

    if "etiketi" in slugs or "етикет" in text:
        return 1408
    if "receipts" in slugs or "касова" in text:
        return 9030
    if "tikso" in slugs or "тиксо" in text:
        return 843
    if "pizza-boxes" in slugs:
        return 11091
    if slugs & {"carton", "plastic", "eco_packages"}:
        return 11085
    if "plastmasovi-pribori" in slugs:
        return 10628
    if "slamki" in slugs:
        return 10088
    if "chuvali-smet" in slugs:
        return 12950
    if slugs & {"vacuum", "plik-cip", "plik-blok", "torbi", "pe-bags"}:
        return 17188
    if slugs & {"stretch-foil", "lldpe-foil-hand", "pvc-foil-hand"}:
        return 17189
    if "alum-foil" in slugs:
        return 17190
    if "kalcuni" in slugs:
        return 31121
    if "maski-za-litse" in slugs:
        return 53765
    if "rykavici" in slugs:
        return 6774
    raise ValueError(f"No Temu category mapping for SKU {source.sku}: {sorted(slugs)}")


def detect_materials(text: str) -> tuple[str, str, str]:
    lowered = text.lower()
    if "nitril" in lowered or "нитрил" in lowered:
        return "Nitrile", "Nitrile", "Nitrile"
    if "non-woven" in lowered or "нетъкан" in lowered:
        return "Non-woven", "Non-woven", "Non-woven"
    if re.search(r"\bpet\b|pет|polyethylene terephthalate", lowered):
        return "PET (polyethylene Terephthalate)", "Polyethylene Terephthalate (PET)", "Polyethylene terephthalate"
    if re.search(r"\bps\b|polystyrene|полистир", lowered):
        return "PS (polystyrene)", "Polystyrene (PS)", "Other Plastic"
    if re.search(r"\bpp\b|polypropylene|полипропилен", lowered):
        return "Polypropylene(PP)", "Polypropylene (PP)", "PP"
    if "hdpe" in lowered:
        return "Plastic", "Polyethylene (PE)", "PE (polyethylene)"
    if "ldpe" in lowered or "lldpe" in lowered or "полиетилен" in lowered:
        return "Plastic", "Polyethylene (PE)", "PE (polyethylene)"
    if "алумини" in lowered or "aluminum" in lowered:
        return "Aluminum Foil", "Aluminum", "Aluminum"
    if any(word in lowered for word in ("картон", "kraft", "крафт", "paper", "харти")):
        return "Paper", "Paper", "Paper"
    if "акрил" in lowered or "acrylic" in lowered:
        return "Acrylic", "PMMA (Acrylic)", "Acrylic"
    return "Plastic", "No Coating", "Other Plastic"


def category_properties(category_id: int, text: str, thickness_um: float) -> dict[str, Any]:
    material, food_material, major_material = detect_materials(text)
    props: dict[str, Any] = {}
    if category_id == 1408:
        props["t_3_Property:12"] = "Thermal Paper" if "термо" in text.lower() else "Paper"
        props["t_3_Property:1081"] = "Thermal Transfer Printing"
    elif category_id == 9030:
        props["t_3_Property:610"] = "48"
    elif category_id in {843, 6774, 53765}:
        if category_id == 843:
            major_material = "PP"
        props["t_3_Property:1920"] = major_material
        if category_id == 53765:
            props["t_3_Property:1117"] = "3 Years+"
    elif category_id == 31121:
        props["t_3_Property:12"] = "Polyethylene (PE)"
        props["t_3_Property:15:1309"] = 100
    elif category_id in {10628, 11085, 11091, 10088, 17188, 17189, 12950}:
        props["t_3_Property:121"] = material
        if category_id in {10628, 11085, 11091, 10088}:
            props["t_3_Property:8319"] = food_material
        if category_id in {11085, 11091, 17188}:
            props["t_3_Property:686"] = "Coated" if any(k in text.lower() for k in ("екстру", "покрит")) else "Uncoated"
        if category_id == 12950:
            props["t_3_Property:7946 - value"] = round(thickness_um, 2)
            props["t_3_Property:7946 - unit"] = "μm"
    return props


def parse_pack_count(text: str) -> int:
    matches = re.findall(r"(?<!\d)(\d{1,5})\s*(?:бр\.?|броя|pcs?\.?|pieces?)", text, flags=re.I)
    return max(1, int(matches[0])) if matches else 1


def parse_thickness_um(text: str, default: float) -> tuple[float, bool]:
    match = re.search(r"(\d+(?:[.,]\d+)?)\s*(?:микрон(?:а)?|мк|µm|μm)", text, flags=re.I)
    return (number(match.group(1), default), False) if match else (default, True)


def parse_dimensions(text: str, source: SourceSku, fallback: float) -> tuple[float, float, float, str, bool]:
    if source.length_cm > 0 and source.width_cm > 0 and source.height_cm > 0:
        dims = (source.length_cm, source.width_cm, source.height_cm)
        return (*dims, " × ".join(f"{x:g}" for x in dims) + " cm", False)

    normalized = text.replace("×", "x").replace("Х", "x").replace("х", "x")
    three = re.search(r"(\d+(?:[.,]\d+)?)\s*[x/]\s*(\d+(?:[.,]\d+)?)\s*[x/]\s*(\d+(?:[.,]\d+)?)\s*(mm|мм|cm|см)\b", normalized, re.I)
    two = re.search(r"(\d+(?:[.,]\d+)?)\s*[x/]\s*(\d+(?:[.,]\d+)?)\s*(mm|мм|cm|см)\b", normalized, re.I)
    circle = re.search(r"(?:ф|ø|диаметър\s*)(\d+(?:[.,]\d+)?)\s*(cm|см|mm|мм)?", normalized, re.I)
    if three:
        values = [number(three.group(i)) for i in (1, 2, 3)]
        if three.group(4).lower() in {"mm", "мм"}:
            values = [v / 10 for v in values]
        return values[0], values[1], values[2], " × ".join(f"{v:g}" for v in values) + " cm", False
    if two:
        values = [number(two.group(i)) for i in (1, 2)]
        if two.group(3).lower() in {"mm", "мм"}:
            values = [v / 10 for v in values]
        return values[0], values[1], fallback, " × ".join(f"{v:g}" for v in values) + " cm", True
    if circle:
        diameter = number(circle.group(1))
        if str(circle.group(2) or "cm").lower() in {"mm", "мм"}:
            diameter /= 10
        return diameter, diameter, fallback, f"Ø{diameter:g} cm", True
    return fallback, fallback, fallback, f"{fallback:g} × {fallback:g} × {fallback:g} cm", True


def product_price(product: dict[str, Any]) -> tuple[float, float | None]:
    prices = product.get("prices") or {}
    minor = int(prices.get("currency_minor_unit", 2) or 2)
    divisor = 10 ** minor
    price = number(prices.get("price")) / divisor
    regular = number(prices.get("regular_price")) / divisor
    if price <= 0:
        raise ValueError("Price is missing or zero")
    return round(price, minor), round(regular, minor) if regular > price else None


def product_quantity(product: dict[str, Any]) -> int:
    stock_class = str((product.get("stock_availability") or {}).get("class", ""))
    if stock_class == "out-of-stock" or not product.get("is_purchasable", True):
        return 0
    maximum = number((product.get("add_to_cart") or {}).get("maximum"), 999)
    return max(0, int(maximum))


def build_row(product: dict[str, Any], source: SourceSku, default_dimension: float, default_weight_g: int, default_thickness_um: float) -> TemuRow:
    name = clean_html(product.get("name")) or source.name
    short_description = clean_html(product.get("short_description"))
    description = clean_html(product.get("description"))
    combined = "\n".join(part for part in (name, short_description, description) if part)
    category_id = pick_category(product, source)
    price, list_price = product_price(product)
    images = []
    for item in product.get("images") or []:
        url = str(item.get("src") or "").strip()
        if url and url not in images:
            images.append(url)
    if not images:
        raise ValueError("No product image")

    notes: list[str] = []
    weight_g = int(round(source.weight_kg * 1000)) if source.weight_kg > 0 else default_weight_g
    if source.weight_kg <= 0:
        notes.append(f"weight fallback={default_weight_g}g")
    length, width, height, dimensions_text, used_dim_fallback = parse_dimensions(combined, source, default_dimension)
    if used_dim_fallback:
        notes.append(f"one or more dimension fallbacks={default_dimension}cm")
    thickness, used_thickness_fallback = parse_thickness_um(combined, default_thickness_um)
    if category_id == 12950 and used_thickness_fallback:
        notes.append(f"thickness fallback={default_thickness_um}μm")

    bullet_candidates = [line.lstrip("• ").strip() for line in short_description.splitlines()]
    bullets = [line for line in bullet_candidates if line and "цената е" not in line.lower()][:6]
    pack_count = parse_pack_count(f"{name}\n{short_description}")
    return TemuRow(
        sku=source.sku,
        name=name,
        description=description or short_description,
        bullets=bullets,
        images=images[:95],
        url=str(product.get("permalink") or source.url),
        category_id=category_id,
        price=price,
        list_price=list_price,
        quantity=product_quantity(product),
        weight_g=max(1, weight_g),
        length_cm=max(0.1, round(length, 2)),
        width_cm=max(0.1, round(width, 2)),
        height_cm=max(0.1, round(height, 2)),
        dimensions_text=dimensions_text,
        pack_count=pack_count,
        properties=category_properties(category_id, combined, thickness),
        notes=notes,
    )


def header_map(ws: Any) -> dict[str, list[int]]:
    result: dict[str, list[int]] = {}
    for col in range(1, ws.max_column + 1):
        key = str(ws.cell(4, col).value or "").strip()
        if key:
            result.setdefault(key, []).append(col)
    return result


def set_first(ws: Any, headers: dict[str, list[int]], row: int, key: str, value: Any) -> None:
    columns = headers.get(key)
    if not columns:
        raise KeyError(f"Template column not found: {key}")
    ws.cell(row, columns[0]).value = value


def write_row(ws: Any, headers: dict[str, list[int]], excel_row: int, item: TemuRow, args: argparse.Namespace) -> None:
    set_first(ws, headers, excel_row, "t_1_Category", item.category_id)
    set_first(ws, headers, excel_row, "t_1_Product type", "Normal product")
    set_first(ws, headers, excel_row, "t_1_Product Name", item.name[:500])
    set_first(ws, headers, excel_row, "t_1_Contribution Goods", item.sku)
    set_first(ws, headers, excel_row, "t_1_Contribution SKU", item.sku)
    set_first(ws, headers, excel_row, "t_1_Update or Add", "Add")
    set_first(ws, headers, excel_row, "t_2_Product Description", item.description[:10000])

    bullet_columns = headers.get("t_2_Bullet Point", [])
    for col, value in zip(bullet_columns, item.bullets):
        ws.cell(excel_row, col).value = value[:500]
    detail_columns = headers.get("t_2_Detail Images URL", [])
    for col, value in zip(detail_columns, item.images):
        ws.cell(excel_row, col).value = value

    for key, value in item.properties.items():
        set_first(ws, headers, excel_row, key, value)

    set_first(ws, headers, excel_row, "t_4_Variation Theme", "Items")
    set_first(ws, headers, excel_row, "t_4_Custom Spec:17020", f"{item.pack_count} pcs")
    sku_image_columns = headers.get("t_6_SKU Images URL", [])
    for col, value in zip(sku_image_columns, item.images[:10]):
        ws.cell(excel_row, col).value = value

    set_first(ws, headers, excel_row, "t_6_Dimensions", item.dimensions_text)
    set_first(ws, headers, excel_row, "t_6_Quantity", item.quantity)
    set_first(ws, headers, excel_row, "t_6_Base Price - EUR", item.price)
    set_first(ws, headers, excel_row, "t_6_Reference Link", item.url)
    if item.list_price is not None:
        set_first(ws, headers, excel_row, "t_6_List Price - EUR", item.list_price)
    else:
        set_first(ws, headers, excel_row, "t_6_Not available for List price", "N/A")
    set_first(ws, headers, excel_row, "t_6_Weight - g", item.weight_g)
    set_first(ws, headers, excel_row, "t_6_Length - cm", item.length_cm)
    set_first(ws, headers, excel_row, "t_6_Width - cm", item.width_cm)
    set_first(ws, headers, excel_row, "t_6_Height - cm", item.height_cm)
    set_first(ws, headers, excel_row, "t_6_SKU type", "Multi-piece set" if item.pack_count > 1 else "Single set")
    set_first(ws, headers, excel_row, "t_6_Individually packed", "No")
    set_first(ws, headers, excel_row, "t_6_Total packaging quantity", 1)
    set_first(ws, headers, excel_row, "t_6_Packaging unit", "pack")
    set_first(ws, headers, excel_row, "t_6_Total item quantity", item.pack_count)
    set_first(ws, headers, excel_row, "t_6_Item unit", "piece")
    set_first(ws, headers, excel_row, "t_7_Shipping Template", args.shipping_template)
    set_first(ws, headers, excel_row, "t_7_Handling Time", args.handling_time)
    set_first(ws, headers, excel_row, "t_7_Fulfillment Channel", "I will ship this item myself")
    set_first(ws, headers, excel_row, "t_7_Item Tax Code", "GEN STANDARD")
    set_first(ws, headers, excel_row, "t_8_Country/Region of Origin", args.country_origin)
    set_first(ws, headers, excel_row, "t_8_Governance Property:1100100115", item.sku)
    set_first(ws, headers, excel_row, "t_8_Governance Property:3", args.manufacturer)
    # EU responsible person is deliberately left blank; it must match a person
    # already configured in the seller account.


def output_parts(template: Path, rows: list[TemuRow], output_dir: Path, max_rows: int, args: argparse.Namespace) -> list[Path]:
    output_dir.mkdir(parents=True, exist_ok=True)
    paths: list[Path] = []
    for part_no, group in enumerate(chunks(rows, max_rows), start=1):
        output_path = output_dir / f"TEMU_TOTALPACK_UPLOAD_part_{part_no:03d}.xlsx"
        shutil.copy2(template, output_path)
        wb = load_workbook(output_path)
        ws = wb["Template"]
        headers = header_map(ws)
        for offset, item in enumerate(group, start=5):
            write_row(ws, headers, offset, item, args)
        try:
            wb.calculation.fullCalcOnLoad = True
            wb.calculation.forceFullCalc = True
        except AttributeError:
            pass
        wb.save(output_path)
        wb.close()
        paths.append(output_path)
    return paths


def write_report(path: Path, records: list[dict[str, Any]]) -> None:
    fields = ["sku", "status", "name", "url", "temu_category", "quantity", "price_eur", "notes"]
    with path.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields)
        writer.writeheader()
        writer.writerows(records)


def parse_args(argv: list[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--template", type=Path, default=Path(DEFAULT_TEMPLATE))
    parser.add_argument("--sku-file", type=Path, default=Path(DEFAULT_SKU_FILE))
    parser.add_argument("--output-dir", type=Path, default=Path("output"))
    parser.add_argument("--limit", type=int, default=0, help="0 = all selected SKUs")
    parser.add_argument("--search", default="", help="Filter SKU/name/category text")
    parser.add_argument("--only-in-stock", action="store_true")
    parser.add_argument("--max-rows", type=int, default=1900)
    parser.add_argument("--timeout", type=int, default=45)
    parser.add_argument("--retries", type=int, default=4)
    parser.add_argument("--default-dimension-cm", type=float, default=10.0)
    parser.add_argument("--default-weight-g", type=int, default=100)
    parser.add_argument("--default-thickness-um", type=float, default=20.0)
    parser.add_argument("--shipping-template", default="boxnow")
    parser.add_argument("--handling-time", default="1 Day")
    parser.add_argument("--country-origin", default="Bulgaria")
    parser.add_argument("--manufacturer", default="ULTRAPAK")
    parser.add_argument("--api-fixture", type=Path, default=None, help=argparse.SUPPRESS)
    parser.add_argument("--log-level", default="INFO", choices=["DEBUG", "INFO", "WARNING", "ERROR"])
    return parser.parse_args(argv)


def main(argv: list[str] | None = None) -> int:
    args = parse_args(argv)
    logging.basicConfig(level=getattr(logging, args.log_level), format="%(levelname)s %(message)s")
    if not args.template.exists():
        raise FileNotFoundError(args.template)
    if not args.sku_file.exists():
        raise FileNotFoundError(args.sku_file)
    if args.max_rows < 1 or args.max_rows > 4996:
        raise ValueError("--max-rows must be between 1 and 4996")

    sources = read_sku_file(args.sku_file)
    if args.search.strip():
        needle = args.search.strip().casefold()
        sources = [s for s in sources if needle in f"{s.sku} {s.name} {s.source_category}".casefold()]
    if args.limit > 0:
        sources = sources[: args.limit]
    if not sources:
        raise RuntimeError("No SKUs selected")

    LOG.info("Selected %s unique SKU(s)", len(sources))
    product_map = fetch_products([s.sku for s in sources], args.timeout, args.retries, args.api_fixture)
    rows: list[TemuRow] = []
    report: list[dict[str, Any]] = []
    for source in sources:
        product = product_map.get(source.sku)
        if not product:
            LOG.warning("SKU %s was not found on totalpack.bg", source.sku)
            report.append({"sku": source.sku, "status": "NOT_FOUND", "name": source.name, "url": source.url, "temu_category": "", "quantity": "", "price_eur": "", "notes": "Exact SKU not returned by the site API"})
            continue
        try:
            row = build_row(product, source, args.default_dimension_cm, args.default_weight_g, args.default_thickness_um)
            if args.only_in_stock and row.quantity <= 0:
                status = "SKIPPED_OUT_OF_STOCK"
            else:
                status = "OK_WITH_FALLBACKS" if row.notes else "OK"
                rows.append(row)
            report.append({"sku": row.sku, "status": status, "name": row.name, "url": row.url, "temu_category": row.category_id, "quantity": row.quantity, "price_eur": f"{row.price:.2f}", "notes": "; ".join(row.notes)})
        except Exception as exc:  # keep processing the remaining requested SKUs
            LOG.exception("Could not process SKU %s", source.sku)
            report.append({"sku": source.sku, "status": "ERROR", "name": source.name, "url": source.url, "temu_category": "", "quantity": "", "price_eur": "", "notes": str(exc)})

    args.output_dir.mkdir(parents=True, exist_ok=True)
    files = output_parts(args.template, rows, args.output_dir, args.max_rows, args) if rows else []
    report_path = args.output_dir / "totalpack_scrape_report.csv"
    write_report(report_path, report)
    summary = {
        "generated_at_utc": datetime.now(timezone.utc).isoformat(),
        "requested_skus": len(sources),
        "written_rows": len(rows),
        "not_found": sum(r["status"] == "NOT_FOUND" for r in report),
        "errors": sum(r["status"] == "ERROR" for r in report),
        "skipped_out_of_stock": sum(r["status"] == "SKIPPED_OUT_OF_STOCK" for r in report),
        "output_files": [p.name for p in files],
    }
    (args.output_dir / "summary.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    LOG.info("Finished: %s rows written to %s Excel file(s)", len(rows), len(files))
    LOG.info("Report: %s", report_path)
    return 0 if rows else 2


if __name__ == "__main__":
    sys.exit(main())
